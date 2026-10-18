import os, re, io, csv, hmac, secrets, sqlite3, smtplib, hashlib, ssl
from email.message import EmailMessage
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from datetime import datetime, timezone
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file, abort
from werkzeug.security import check_password_hash
import pdfplumber
from openpyxl import load_workbook
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', secrets.token_hex(32))
app.config.update(MAX_CONTENT_LENGTH=12*1024*1024, SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax', SESSION_COOKIE_SECURE=os.environ.get('HTTPS_ONLY','0')=='1')
DB=os.environ.get('DATABASE_PATH', os.path.join(os.path.dirname(__file__),'payroll.sqlite3'))
DRIVERS={'BRAKLE':'','FLOAND':'01097859','ANADUR':'01337496','WALKEV':'','CERAEX':'','SULLED':'','MORTAR5':'','MARTMI':'01214498','WILSPA':'01368410','FONCES7':'','OWESTA1':'','TARMOR':'01404076'}

def db():
    c=sqlite3.connect(DB); c.row_factory=sqlite3.Row; return c

def init():
    with db() as c:
        c.execute('CREATE TABLE IF NOT EXISTS drivers (id TEXT PRIMARY KEY, transponder TEXT NOT NULL DEFAULT "")')
        c.execute('CREATE TABLE IF NOT EXISTS email_log (week TEXT NOT NULL, driver TEXT NOT NULL, email TEXT NOT NULL, fingerprint TEXT NOT NULL, sent_at TEXT NOT NULL, PRIMARY KEY(week,driver,fingerprint,email))')
        cols=[r[1] for r in c.execute('PRAGMA table_info(drivers)')]
        if 'email' not in cols: c.execute('ALTER TABLE drivers ADD COLUMN email TEXT NOT NULL DEFAULT ""')
        if 'rate_bp' not in cols: c.execute('ALTER TABLE drivers ADD COLUMN rate_bp INTEGER NOT NULL DEFAULT 5000')
        c.execute('CREATE TABLE IF NOT EXISTS payroll (week TEXT NOT NULL, driver TEXT NOT NULL, gross_cents INTEGER NOT NULL, toll_cents INTEGER NOT NULL, PRIMARY KEY(week,driver))')
        for k,v in DRIVERS.items(): c.execute('INSERT OR IGNORE INTO drivers(id,transponder) VALUES (?,?)',(k,v))
        pcols=[r[1] for r in c.execute('PRAGMA table_info(payroll)')]
        if 'rate_bp' not in pcols: c.execute('ALTER TABLE payroll ADD COLUMN rate_bp INTEGER NOT NULL DEFAULT 5000')
init()

def login_required(f):
    @wraps(f)
    def wrapped(*a,**kw):
        if not session.get('auth'): return redirect(url_for('login'))
        return f(*a,**kw)
    return wrapped

@app.before_request
def csrf():
    if 'csrf' not in session: session['csrf']=secrets.token_urlsafe(32)
    if request.method=='POST' and not hmac.compare_digest(request.form.get('csrf',''),session['csrf']): abort(400)

@app.context_processor
def context(): return {'csrf':session.get('csrf','')}

def cents(v):
    try:
        d=Decimal(str(v).replace('$','').replace(',','').strip())
        if not d.is_finite() or abs(d)>100000000: raise ValueError()
        return int((d*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
    except (InvalidOperation,ValueError,TypeError): raise ValueError('Enter a valid dollar amount.')

def pay_amount(gross,toll,rate_bp):
    return int((Decimal(gross-toll)*Decimal(rate_bp)/Decimal(10000)).quantize(Decimal('1'),rounding=ROUND_HALF_UP))

def money(c): return f'${c/100:,.2f}'
app.jinja_env.filters['money']=money

def drivers():
    with db() as c: return c.execute('SELECT * FROM drivers ORDER BY id').fetchall()

def weeks():
    with db() as c: return c.execute('SELECT DISTINCT week FROM payroll ORDER BY week DESC').fetchall()

@app.route('/login',methods=['GET','POST'])
def login():
    if request.method=='POST':
        configured=os.environ.get('ADMIN_PASSWORD_HASH')
        if not configured:
            flash('Set ADMIN_PASSWORD_HASH before using the app.'); return render_template('login.html')
        if check_password_hash(configured,request.form.get('password','')):
            session.clear(); session['auth']=True; session['csrf']=secrets.token_urlsafe(32)
            return redirect(url_for('index'))
        flash('Incorrect password.')
    return render_template('login.html')

@app.post('/logout')
@login_required
def logout(): session.clear(); return redirect(url_for('login'))

@app.route('/')
@login_required
def index():
    week=request.args.get('week') or datetime.now().date().isoformat()
    with db() as c:
        rows=c.execute('SELECT d.id,d.transponder,d.rate_bp AS default_rate,COALESCE(p.rate_bp,d.rate_bp) rate_bp,COALESCE(p.gross_cents,0) gross,COALESCE(p.toll_cents,0) toll FROM drivers d LEFT JOIN payroll p ON d.id=p.driver AND p.week=? ORDER BY d.id',(week,)).fetchall()
    items=[dict(r)|{'pay':pay_amount(r['gross'],r['toll'],r['rate_bp'])} for r in rows]
    totals={k:sum(x[k] for x in items) for k in ('gross','toll','pay')}
    return render_template('index.html',week=week,items=items,totals=totals,weeks=weeks(),smtp_ready=all(os.environ.get(k) for k in ('SMTP_HOST','SMTP_USER','SMTP_PASSWORD','SMTP_FROM')))

@app.post('/save')
@login_required
def save():
    week=request.form.get('week','')
    try: datetime.strptime(week,'%Y-%m-%d')
    except ValueError: abort(400)
    with db() as c:
        for d in drivers():
            k=d['id']
            try: gross=cents(request.form.get('gross_'+k,'0')); toll=cents(request.form.get('toll_'+k,'0'))
            except ValueError: flash('Invalid amount for '+k); return redirect(url_for('index',week=week))
            if gross<0 or toll<0: flash('Gross and tolls must be nonnegative.'); return redirect(url_for('index',week=week))
            c.execute('INSERT INTO payroll(week,driver,gross_cents,toll_cents,rate_bp) VALUES (?,?,?,?,?) ON CONFLICT(week,driver) DO UPDATE SET gross_cents=excluded.gross_cents,toll_cents=excluded.toll_cents',(week,k,gross,toll,d['rate_bp']))
    flash('Payroll saved. Review totals before paying drivers.')
    return redirect(url_for('index',week=week))

@app.post('/driver/rate')
@login_required
def driver_rate():
    k=request.form.get('id','').strip().upper()
    try:
        rate_bp=int((Decimal(request.form.get('rate',''))*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
        if not 0<=rate_bp<=10000: raise ValueError()
    except (ValueError,InvalidOperation):
        flash('Invalid percentage.'); return redirect(url_for('index'))
    with db() as c:
        c.execute('UPDATE drivers SET rate_bp=? WHERE id=?',(rate_bp,k))
    flash('Future payroll rate updated. Previously saved weeks remain unchanged.')
    return redirect(url_for('index'))

@app.post('/driver')
@login_required
def driver():
    k=request.form.get('id','').strip().upper(); t=request.form.get('transponder','').strip(); email=request.form.get('email','').strip().lower(); rate=request.form.get('rate','50').strip()
    try:
        rate_bp=int((Decimal(rate)*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
        if not 0<=rate_bp<=10000: raise ValueError()
    except (ValueError,InvalidOperation):
        flash('Pay percentage must be between 0 and 100.'); return redirect(url_for('index'))
    if not re.fullmatch(r'[A-Z0-9_-]{2,25}',k) or (t and not re.fullmatch(r'\d{1,20}',t)) or (email and not re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+',email)):
        flash('Invalid driver ID or transponder.'); return redirect(url_for('index'))
    with db() as c: c.execute('INSERT INTO drivers(id,transponder,email,rate_bp) VALUES (?,?,?,?) ON CONFLICT(id) DO UPDATE SET transponder=excluded.transponder,email=excluded.email,rate_bp=excluded.rate_bp',(k,t,email,rate_bp))
    flash('Driver saved.'); return redirect(url_for('index'))

# PDF extraction is deliberately a reviewable suggestion, not an automatic payroll posting.
@app.post('/inspect/pdf')
@login_required
def inspect_pdf():
    f=request.files.get('file')
    if not f or not f.filename.lower().endswith('.pdf'): abort(400)
    try:
        with pdfplumber.open(io.BytesIO(f.read())) as pdf:
            lines='\n'.join((p.extract_text() or '') for p in pdf.pages[:30]).splitlines()
        found=[]
        for line in lines:
            for d in drivers():
                if re.search(r'(?<!\w)'+re.escape(d['id'])+r'(?!\w)',line,re.I):
                    found.append((d['id'],line[:240]))
        return render_template('inspect.html',kind='Settlement PDF',found=found[:100],note='Amounts in the PDF have NOT been imported. Check which column represents gross revenue, then enter it on the payroll screen. Send a sample report to configure reliable automatic parsing.')
    except Exception: flash('Could not read this PDF. Scanned/image-only PDFs require OCR.'); return redirect(url_for('index'))

@app.post('/inspect/tolls')
@login_required
def inspect_tolls():
    f=request.files.get('file')
    if not f or not f.filename.lower().endswith(('.xlsx','.xlsm')): abort(400)
    try:
        wb=load_workbook(io.BytesIO(f.read()),read_only=True,data_only=True)
        mapping={d['transponder']:d['id'] for d in drivers() if d['transponder']}
        matches={k:[] for k in mapping.values()}
        preview=[]
        for sh in wb:
            for row in sh.iter_rows(min_row=1,max_row=15000,values_only=True):
                vals=[str(x).strip() if x is not None else '' for x in row]
                for trans,k in mapping.items():
                    if any(v.lstrip('0')==trans.lstrip('0') for v in vals if v.isdigit()):
                        if len(matches[k])<20: matches[k].append(' | '.join(vals[:12])[:230])
        for k,lines in matches.items():
            for line in lines: preview.append((k,line))
        wb.close()
        return render_template('inspect.html',kind='Toll Excel',found=preview[:100],note='Matching transponder rows are shown for verification. Tolls are NOT deducted automatically until the amount/date columns are confirmed using a sample file.')
    except Exception: flash('Could not read the Excel file. Upload an .xlsx file.'); return redirect(url_for('index'))


# Import both source reports, preview exact figures, and require approval before posting.
@app.post('/import/preview')
@login_required
def import_preview():
    pdf_file=request.files.get('settlement'); excel_file=request.files.get('tolls')
    week=request.form.get('week','')
    try: datetime.strptime(week,'%Y-%m-%d')
    except ValueError: abort(400)
    if not pdf_file or not excel_file or not pdf_file.filename.lower().endswith('.pdf') or not excel_file.filename.lower().endswith('.xlsx'): abort(400)
    try:
        with pdfplumber.open(io.BytesIO(pdf_file.read())) as pdf:
            source='\n'.join(p.extract_text() or '' for p in pdf.pages)
        sections=list(re.finditer(r'Settlement sheet for Driver:([A-Z0-9]+)',source))
        gross={}
        for i,m in enumerate(sections):
            part=source[m.end():sections[i+1].start() if i+1<len(sections) else len(source)]
            match=re.search(r'Payments Total:\s*([\d,]+\.\d{2})',part)
            if match: gross[m.group(1)]=cents(match.group(1))
        if not gross: raise ValueError('No driver payment totals found in PDF.')
        book=load_workbook(io.BytesIO(excel_file.read()),read_only=True,data_only=True)
        trans={str(d['transponder']).lstrip('0'):d['id'] for d in drivers() if d['transponder']}
        tolls={k:0 for k in gross}; unknown={}; counted=0
        for sh in book:
            rows=sh.iter_rows(values_only=True)
            header=next(rows,None)
            if not header: continue
            cols={str(v).strip().upper():i for i,v in enumerate(header) if v is not None}
            if not {'TRANSPONDER NUMBER','AMOUNT','TRANSACTION TYPE'}.issubset(cols): continue
            for row in rows:
                def val(key):
                    i=cols[key]; return row[i] if i<len(row) else None
                typ=str(val('TRANSACTION TYPE') or '').upper()
                if 'TOLL' not in typ: continue
                raw=str(val('TRANSPONDER NUMBER') or '').strip()
                if not raw.isdigit(): continue
                amt=str(val('AMOUNT') or '').strip().replace('(','').replace(')','').replace('$','').replace(',','').replace('-','')
                if not amt: continue
                charge=abs(cents(amt))
                key=raw.lstrip('0'); driver=trans.get(key)
                if driver and driver in tolls: tolls[driver]+=charge; counted+=1
                else: unknown[raw]=unknown.get(raw,0)+charge
        book.close()
        rate_map={d['id']:d['rate_bp'] for d in drivers()}
        rows=[{'id':k,'gross':v,'toll':tolls.get(k,0),'rate_bp':rate_map.get(k,5000),'pay':pay_amount(v,tolls.get(k,0),rate_map.get(k,5000))} for k,v in gross.items()]
        session['pending_import']={'week':week,'rows':rows}
        return render_template('import_preview.html',week=week,rows=rows,unknown=unknown,counted=counted)
    except Exception as exc:
        flash('Import failed: '+str(exc)[:220]); return redirect(url_for('index',week=week))

@app.post('/import/confirm')
@login_required
def import_confirm():
    pending=session.pop('pending_import',None)
    if not pending or request.form.get('confirm')!='yes': abort(400)
    week=pending['week']
    with db() as c:
        for r in pending['rows']:
            c.execute('INSERT INTO payroll (week,driver,gross_cents,toll_cents,rate_bp) VALUES (?,?,?,?,?) ON CONFLICT(week,driver) DO UPDATE SET gross_cents=excluded.gross_cents,toll_cents=excluded.toll_cents,rate_bp=excluded.rate_bp',(week,r['id'],r['gross'],r['toll'],r['rate_bp']))
    flash('Imported figures saved. Verify all rows before emailing or paying drivers.')
    return redirect(url_for('index',week=week))

@app.get('/export.csv')
@login_required
def export_csv():
    week=request.args.get('week','')
    with db() as c: rows=c.execute('SELECT driver,gross_cents,toll_cents,rate_bp FROM payroll WHERE week=? ORDER BY driver',(week,)).fetchall()
    output=io.StringIO(); w=csv.writer(output); w.writerow(['Week','Driver','Gross','Tolls','Driver Pay Rate','Driver Pay','Company Share'])
    for r in rows:
        net=r['gross_cents']-r['toll_cents']; pay=pay_amount(r['gross_cents'],r['toll_cents'],r['rate_bp'])
        w.writerow([week,r['driver'],f'{r["gross_cents"]/100:.2f}',f'{r["toll_cents"]/100:.2f}',f'{r["rate_bp"]/100:.2f}%',f'{pay/100:.2f}',f'{(net-pay)/100:.2f}'])
    b=io.BytesIO(output.getvalue().encode('utf-8-sig'))
    return send_file(b,mimetype='text/csv',as_attachment=True,download_name=f'CERA_EXPRESS_{week}_payroll.csv')

@app.get('/statement.pdf')
@login_required
def statement():
    week=request.args.get('week',''); driver=request.args.get('driver','')
    with db() as c: r=c.execute('SELECT * FROM payroll WHERE week=? AND driver=?',(week,driver)).fetchone()
    if not r: abort(404)
    net=r['gross_cents']-r['toll_cents']; pay=pay_amount(r['gross_cents'],r['toll_cents'],r['rate_bp'])
    buf=build_statement(week,driver,r)
    return send_file(buf,mimetype='application/pdf',as_attachment=True,download_name=f'{driver}_{week}_settlement.pdf')

def build_statement(week,driver,r):
    net=r['gross_cents']-r['toll_cents']; pay=pay_amount(r['gross_cents'],r['toll_cents'],r['rate_bp'])
    buf=io.BytesIO(); doc=SimpleDocTemplate(buf); styles=getSampleStyleSheet()
    story=[Paragraph('CERA EXPRESS LLC',styles['Title']),Paragraph('Driver settlement statement',styles['Heading2']),Spacer(1,15)]
    data=[['Week',week],['Driver',driver],['Gross revenue',money(r['gross_cents'])],['Tolls',money(r['toll_cents'])],['Revenue after tolls',money(net)],[f'Driver pay ({r["rate_bp"]/100:g}%)',money(pay)],['Company share',money(net-pay)]]
    t=Table(data,colWidths=[210,220]); t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.5,colors.lightgrey),('PADDING',(0,0),(-1,-1),10),('BACKGROUND',(0,-2),(-1,-2),colors.whitesmoke)])); story.append(t)
    story.append(Spacer(1,15)); story.append(Paragraph('Review and approve before payment. Other deductions and reimbursements are not included.',styles['Normal']))
    doc.build(story); buf.seek(0)
    return buf

def fingerprint(row):
    return hashlib.sha256(f"{row['week']}|{row['driver']}|{row['gross_cents']}|{row['toll_cents']}|{row['rate_bp']}".encode()).hexdigest()

@app.get('/email/preview')
@login_required
def email_preview():
    week=request.args.get('week','')
    with db() as c:
        rows=c.execute('SELECT p.*,d.email FROM payroll p JOIN drivers d ON d.id=p.driver WHERE p.week=? ORDER BY p.driver',(week,)).fetchall()
        entries=[]
        for r in rows:
            sent=c.execute('SELECT sent_at FROM email_log WHERE week=? AND driver=? AND email=? AND fingerprint=?',(week,r['driver'],r['email'],fingerprint(r))).fetchone() if r['email'] else None
            entries.append({'driver':r['driver'],'email':r['email'],'gross':r['gross_cents'],'toll':r['toll_cents'],'sent':sent['sent_at'] if sent else None})
    return render_template('email_preview.html',week=week,entries=entries,smtp_ready=all(os.environ.get(k) for k in ('SMTP_HOST','SMTP_USER','SMTP_PASSWORD','SMTP_FROM')))

@app.post('/email/send')
@login_required
def email_send():
    week=request.form.get('week','')
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}',week): abort(400)
    if request.form.get('confirm')!='yes': flash('You must confirm the payroll was reviewed.'); return redirect(url_for('email_preview',week=week))
    required=('SMTP_HOST','SMTP_USER','SMTP_PASSWORD','SMTP_FROM')
    if not all(os.environ.get(k) for k in required): flash('Email is not configured on the server.'); return redirect(url_for('email_preview',week=week))
    selected=request.form.getlist('drivers')
    if not selected: flash('Select at least one driver.'); return redirect(url_for('email_preview',week=week))
    if len(selected)>150: abort(400)
    sent_count=0; skipped=0; failed=[]
    for driver_id in dict.fromkeys(selected):
        with db() as c:
            r=c.execute('SELECT p.*,d.email FROM payroll p JOIN drivers d ON d.id=p.driver WHERE p.week=? AND p.driver=?',(week,driver_id)).fetchone()
            if not r or not re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+',r['email'] or ''): skipped+=1; continue
            fp=fingerprint(r)
            if c.execute('SELECT 1 FROM email_log WHERE week=? AND driver=? AND fingerprint=? AND email=?',(week,driver_id,fp,r['email'])).fetchone(): skipped+=1; continue
        net=r['gross_cents']-r['toll_cents']; pay=pay_amount(r['gross_cents'],r['toll_cents'],r['rate_bp'])
        msg=EmailMessage(); msg['Subject']=f'CERA EXPRESS LLC | Settlement {week} | {driver_id}'
        msg['From']=os.environ['SMTP_FROM']; msg['To']=r['email']
        msg.set_content(f'Hello,\n\nAttached is your CERA EXPRESS LLC settlement for the week of {week}.\nDriver: {driver_id}\nGross: {money(r["gross_cents"])}\nTolls: {money(r["toll_cents"])}\nDriver pay ({r["rate_bp"]/100:g}% after tolls): {money(pay)}\n\nPlease contact dispatch if you have questions.\n\nCERA EXPRESS LLC')
        msg.add_attachment(build_statement(week,driver_id,r).getvalue(),maintype='application',subtype='pdf',filename=f'{driver_id}_{week}_settlement.pdf')
        try:
            port=int(os.environ.get('SMTP_PORT','587'))
            if os.environ.get('SMTP_SSL','0')=='1':
                server=smtplib.SMTP_SSL(os.environ['SMTP_HOST'],port,timeout=20,context=ssl.create_default_context())
            else:
                server=smtplib.SMTP(os.environ['SMTP_HOST'],port,timeout=20)
            with server:
                if os.environ.get('SMTP_SSL','0')!='1': server.starttls(context=ssl.create_default_context())
                server.login(os.environ['SMTP_USER'],os.environ['SMTP_PASSWORD'])
                server.send_message(msg)
            with db() as c: c.execute('INSERT OR IGNORE INTO email_log VALUES (?,?,?,?,?)',(week,driver_id,r['email'],fp,datetime.now(timezone.utc).isoformat()))
            sent_count+=1
        except Exception:
            failed.append(driver_id)
    flash(f'Emails sent: {sent_count}. Skipped (missing email/already sent): {skipped}. Failed: {", ".join(failed) if failed else "none"}. Check mail provider delivery logs for final delivery status.')
    return redirect(url_for('email_preview',week=week))

if __name__=='__main__': app.run(debug=False)
